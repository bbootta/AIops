"""1슬라이드 PPTX 여러 개를 순서대로 한 파일로 병합한다.

슬라이드 파트와 그 하위 파트(레이아웃, 마스터, 테마, 이미지, 차트, 임베디드 파일)를
OPC 패키지 단위로 복사한다. 같은 템플릿에서 나온 레이아웃·마스터는 내용 해시로
중복을 제거한다. 발표노트는 텍스트로 옮긴다.
"""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PartFactory, _Relationship
from pptx.opc.packuri import PackURI

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SKIP_RELS = {RT.NOTES_SLIDE}
DEDUP_SUFFIXES = (".slideLayout+xml", ".slideMaster+xml")  # 테마는 마스터마다 따로 둔다


class MergeError(ValueError):
    pass


def check_single(path: Path, expect_size: tuple[int, int] | None = None) -> tuple[int, int]:
    """1슬라이드 여부와 화면 크기를 확인하고 (가로, 세로) EMU를 반환한다."""
    prs = Presentation(str(path))
    if len(prs.slides) != 1:
        raise MergeError(f"{path.name}: 슬라이드 {len(prs.slides)}장 (1장이어야 함)")
    size = (prs.slide_width, prs.slide_height)
    if expect_size and size != expect_size:
        raise MergeError(f"{path.name}: 화면 크기 {size}가 기준 {expect_size}와 다름")
    return size


def _key(part) -> str:
    """템플릿 파트의 동일성 키. 레이아웃은 마스터까지, 마스터는 테마까지 포함해 비교한다."""
    h = hashlib.sha256(part.blob)
    if part.content_type.endswith(".slideLayout+xml"):
        h.update(_key(part.part_related_by(RT.SLIDE_MASTER)).encode())
    elif part.content_type.endswith(".slideMaster+xml"):
        h.update(_key(part.part_related_by(RT.THEME)).encode())
    return h.hexdigest()


class _Copier:
    def __init__(self, dst_prs):
        self.prs = dst_prs
        self.pkg = dst_prs.part.package
        self.used = {str(p.partname) for p in self.pkg.iter_parts()}
        # 동일성 키 -> 대상 파트 (같은 템플릿의 레이아웃·마스터 재사용)
        self.by_hash = {
            _key(p): p for p in self.pkg.iter_parts()
            if p.content_type.endswith(DEDUP_SUFFIXES)
        }
        self.ids = self._layout_master_ids()

    def _layout_master_ids(self) -> set[int]:
        ids = set()
        for m in self.prs.slide_masters:
            ids.update(int(e.get("id")) for e in m._element.iter(f"{{{NS_P}}}sldLayoutId"))
        ids.update(int(e.get("id")) for e in self.prs.part._element.iter(f"{{{NS_P}}}sldMasterId"))
        return ids

    def _next_id(self) -> int:
        n = max(self.ids | {2147483647}) + 1
        self.ids.add(n)
        return n

    def _partname(self, src_name: str) -> PackURI:
        stem, ext = src_name.rsplit(".", 1)
        base = re.sub(r"\d+$", "", stem)
        self.used |= {str(p.partname) for p in self.pkg.iter_parts()}  # python-pptx가 만든 파트 포함
        n = 1
        while f"{base}{n}.{ext}" in self.used:
            n += 1
        name = f"{base}{n}.{ext}"
        self.used.add(name)
        return PackURI(name)

    def copy(self, src, memo: dict) -> object:
        if id(src) in memo:
            return memo[id(src)]
        if src.content_type.endswith(DEDUP_SUFFIXES):
            hit = self.by_hash.get(_key(src))
            if hit is not None:
                memo[id(src)] = hit
                return hit
        is_master = src.content_type.endswith(".slideMaster+xml")
        blob = src.blob
        if is_master:  # 레이아웃 ID는 파일 전체에서 유일해야 한다
            root = etree.fromstring(blob)
            for e in root.iter(f"{{{NS_P}}}sldLayoutId"):
                e.set("id", str(self._next_id()))
            blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        dst = PartFactory(self._partname(str(src.partname)), src.content_type, self.pkg, blob)
        memo[id(src)] = dst
        if src.content_type.endswith(DEDUP_SUFFIXES):
            self.by_hash[_key(src)] = dst
        for rId, rel in src.rels.items():
            if rel.reltype in SKIP_RELS:
                continue
            target = rel.target_ref if rel.is_external else self.copy(rel.target_part, memo)
            dst.rels._rels[rId] = _Relationship(
                dst.partname.baseURI, rId, rel.reltype, rel._target_mode, target)
        if is_master:
            rId = self.prs.part.relate_to(dst, RT.SLIDE_MASTER)
            lst = self.prs.part._element.find(f"{{{NS_P}}}sldMasterIdLst")
            e = etree.SubElement(lst, f"{{{NS_P}}}sldMasterId")
            e.set("id", str(self._next_id()))
            e.set(f"{{{NS_R}}}id", rId)
        return dst


def _drop_notes_rels(slide_part) -> None:
    for rId in [r for r, rel in slide_part.rels.items() if rel.reltype == RT.NOTES_SLIDE]:
        slide_part.rels.pop(rId)


def merge(paths: list[Path], out: Path) -> int:
    """`paths` 순서대로 병합해 `out`에 저장하고 슬라이드 수를 반환한다."""
    if not paths:
        raise MergeError("병합할 파일이 없음")
    size = check_single(paths[0])
    for p in paths[1:]:
        check_single(p, size)
    dst = Presentation(str(paths[0]))
    copier = _Copier(dst)
    for p in paths[1:]:
        src = Presentation(str(p))
        s_slide = src.slides[0]
        new_part = copier.copy(s_slide.part, {})
        _drop_notes_rels(new_part)
        rId = dst.part.relate_to(new_part, RT.SLIDE)
        dst.slides._sldIdLst.add_sldId(rId)
        if s_slide.has_notes_slide:
            text = s_slide.notes_slide.notes_text_frame.text
            if text:
                dst.slides[-1].notes_slide.notes_text_frame.text = text
    dst.save(str(out))
    n = len(Presentation(str(out)).slides)
    if n != len(paths):
        raise MergeError(f"병합 결과 {n}장, 입력 {len(paths)}장")
    return n


def to_pdf(pptx_path: Path, outdir: Path) -> Path:
    """LibreOffice로 PDF를 만든다."""
    import subprocess

    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(pptx_path)],
        check=True, capture_output=True, timeout=300)
    pdf = outdir / (pptx_path.stem + ".pdf")
    if not pdf.exists():
        raise MergeError("PDF 변환 실패")
    return pdf


__all__ = ["MergeError", "check_single", "merge", "to_pdf"]
