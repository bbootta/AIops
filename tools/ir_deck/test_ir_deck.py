"""디캠프 9기 IR 덱 실행기: 버전 레지스트리, PPTX 병합, Slack 처리 흐름."""
import hashlib
import re
import sys
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from ir_deck import run  # noqa: E402
from ir_deck.merge import MergeError, merge  # noqa: E402
from ir_deck.registry import SLIDE_IDS, Registry, is_build_command  # noqa: E402
from ir_deck.slack import Slack  # noqa: E402

FN = "OLA_DCAMP9_S09_BusinessModel_v0.1.0_20261004_JJ.pptx"


def _slide(path, title, layout=1, chart=False, notes=None, width=Inches(13.333), theme_tag=None):
    p = Presentation()
    p.slide_width, p.slide_height = width, Inches(7.5)
    s = p.slides.add_slide(p.slide_layouts[layout])
    s.shapes.title.text = title
    if chart:
        cd = CategoryChartData()
        cd.categories = ["A", "B"]
        cd.add_series("S", (1, 2))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(2), Inches(4), Inches(3), cd)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    if theme_tag:  # 다른 템플릿을 흉내 내려고 마스터 내용을 바꾼다
        p.slide_masters[0].name = theme_tag
    p.save(path)
    return Path(path)


# ---------- 레지스트리 ----------

def test_filename_version_and_auto_patch():
    r = Registry()
    assert "S09 v0.1.0 초안" in r.register_upload(FN, "F1", "a", "U1", "1")
    msg = r.register_upload("OLA_DCAMP9_S09_BM.pptx", "F2", "b", "U1", "2")
    assert "v0.1.1" in msg and "자동 부여" in msg
    assert r.register_upload("deck.pptx", "F3", "c", "U1", "3") == ""


def test_same_version_different_content_rejected():
    r = Registry()
    r.register_upload(FN, "F1", "a", "U1", "1")
    assert "거부" in r.register_upload(FN, "F2", "b", "U1", "2")
    assert "같아" in r.register_upload(FN, "F3", "a", "U1", "3")
    assert len(r.versions) == 1


def test_four_eyes_and_adopt_requires_approval():
    r = Registry()
    r.register_upload(FN, "F1", "a", "U1", "1")
    assert "채택할 수" in r.apply_command("채택 S09 v0.1.0", "U2", "2")
    assert "본인" in r.apply_command("승인 S09 v0.1.0", "U1", "3")
    assert "승인 기록" in r.apply_command("승인 S09 v0.1.0", "U2", "4")
    assert "채택" in r.apply_command("채택 S09 v0.1.0", "U2", "5")
    assert r.adopted("S09").reviewer == "U2"
    r.register_upload(FN.replace("v0.1.0", "v0.2.0"), "F2", "b", "U1", "6")
    r.apply_command("승인 S09 v0.2.0", "U3", "7")
    r.apply_command("채택 S09 v0.2.0", "U3", "8")
    assert r.find("S09", "0.1.0").status == "대체"
    assert r.adopted("S09").version == "0.2.0"


def test_plan_fail_closed_and_draft():
    r = Registry()
    r.register_upload(FN, "F1", "a", "U1", "1")
    picks, problems = r.plan(draft=False)
    assert len(problems) == 20 and not picks
    picks, problems = r.plan(draft=True)
    assert [v.slide_id for v in picks] == ["S09"] and not problems
    assert is_build_command("취합") == (True, False)
    assert is_build_command(" 취합 초안 ") == (True, True)
    assert is_build_command("취합해줘") == (False, False)


def test_registry_roundtrip(tmp_path):
    r = Registry()
    r.register_upload(FN, "F1", "a", "U1", "1", link="https://x")
    r.save(tmp_path / "r.json")
    assert Registry.load(tmp_path / "r.json").versions[0].link == "https://x"


# ---------- 병합 ----------

def test_merge_keeps_order_media_chart_notes_and_masters(tmp_path):
    a = _slide(tmp_path / "a.pptx", "S01", notes="노트1")
    b = _slide(tmp_path / "b.pptx", "S02", layout=5, chart=True, notes="노트2")
    c = _slide(tmp_path / "c.pptx", "S03", theme_tag="다른 템플릿")
    out = tmp_path / "out.pptx"
    assert merge([a, b, c], out) == 3
    p = Presentation(str(out))
    assert [s.shapes.title.text for s in p.slides] == ["S01", "S02", "S03"]
    assert any(sh.has_chart for sh in p.slides[1].shapes)
    assert p.slides[1].notes_slide.notes_text_frame.text == "노트2"
    assert len(p.slide_masters) == 2  # 같은 템플릿은 재사용, 다른 템플릿만 추가
    ids = [int(e.get("id")) for m in p.slide_masters for e in m._element.iter(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}sldLayoutId")]
    assert len(ids) == len(set(ids))
    import zipfile
    names = zipfile.ZipFile(out).namelist()
    assert len(names) == len(set(names))  # 파트 이름 중복 없음


def test_merge_rejects_multi_slide_and_size_mismatch(tmp_path):
    a = _slide(tmp_path / "a.pptx", "S01")
    wide43 = _slide(tmp_path / "b.pptx", "S02", width=Inches(10))
    with pytest.raises(MergeError, match="화면 크기"):
        merge([a, wide43], tmp_path / "o.pptx")
    two = Presentation(str(a))
    two.slides.add_slide(two.slide_layouts[0])
    two.save(str(tmp_path / "two.pptx"))
    with pytest.raises(MergeError, match="1장"):
        merge([tmp_path / "two.pptx"], tmp_path / "o.pptx")


# ---------- Slack 처리 흐름 ----------

class FakeSlack:
    """채널 메시지·파일·캔버스를 메모리로 흉내 낸다. 채널 없이 올린 파일은 봇 비공개 파일(`private`)."""

    def __init__(self, messages, files):
        self.messages, self.files = messages, files
        self.posts, self.uploads, self.private, self.canvases = [], [], {}, {}

    def new_messages(self, channel, after):
        return [m for m in self.messages if float(m["ts"]) > float(after)]

    def info(self, file_id):
        return {"permalink": f"https://files/{file_id}"}

    def download(self, file_id):
        if file_id in self.private:
            return {}, self.private[file_id]
        return {"permalink": f"https://files/{file_id}"}, self.files[file_id]

    def upload(self, channel, thread_ts, filename, data, comment=""):
        fid = f"FU{len(self.uploads) + 1}"
        self.uploads.append(filename)
        if channel is None:
            self.private[fid] = data
        return fid, f"https://files/{filename}"

    def post(self, channel, text, thread_ts=None):
        self.posts.append(text)

    def name(self, user_id):
        return user_id

    def form(self, method, **p):
        if method == "auth.test":
            return {"user_id": "UBOT", "bot_id": "BME"}
        if method == "files.list":
            return {"files": [{"id": i, "name": run.STATE_NAME, "created": int(i[2:])} for i in self.private]}
        assert method == "files.delete"
        del self.private[p["file"]]
        return {}

    def call(self, method, **p):
        assert method == "canvases.access.set"
        return {}

    def canvas_create(self, title, md):
        self.canvases["FC"] = md
        return "FC"

    def canvas_replace_all(self, canvas_id, md):
        self.canvases[canvas_id] = md

    def canvas_replace_section(self, canvas_id, text, md):
        self.canvases[canvas_id] = md
        return True


def test_process_end_to_end_draft_build(tmp_path, monkeypatch):
    files, msgs = {}, []
    for i, sid in enumerate(["S01", "S02"], start=1):
        path = _slide(tmp_path / f"{sid}.pptx", sid)
        files[f"F{i}"] = path.read_bytes()
        msgs.append({"ts": f"{i}.0", "user": "U1", "text": "",
                     "files": [{"id": f"F{i}", "name": f"OLA_DCAMP9_{sid}_X_v0.1.0.pptx"}]})
    msgs += [{"ts": "3.0", "user": "U1", "text": "승인 S01 v0.1.0"},
             {"ts": "4.0", "user": "U2", "text": "승인 S01 v0.1.0"},
             {"ts": "5.0", "user": "U2", "text": "취합"},
             {"ts": "6.0", "user": "U2", "text": "취합 초안"}]
    monkeypatch.setattr(run, "to_pdf", lambda p, d: p)  # LibreOffice 없이 흐름만 확인
    monkeypatch.setattr(run, "PDF_READY", True)
    sl, reg = FakeSlack(msgs, files), Registry()
    assert run.process(sl, reg) == (True, False)
    assert reg.cursor == "6.0"
    assert any("본인" in p for p in sl.posts)
    assert reg.builds[0].result.startswith("중단") and "S02 채택본 없음" in reg.builds[0].result
    assert reg.builds[1].result == "성공" and reg.builds[1].pages == 2
    assert len(sl.uploads) == 2 and "DRAFT" in sl.uploads[0]
    assert reg.find("S01", "0.1.0").sha256 == hashlib.sha256(files["F1"]).hexdigest()
    md = run.render_status(reg, lambda u: u)
    assert "최신 취합본: **" in md and all(s in md for s in SLIDE_IDS)
    assert "—" not in md and "–" not in md


def test_process_skips_self_limits_bots_and_isolates_errors(tmp_path, monkeypatch):
    s01 = _slide(tmp_path / "S01.pptx", "S01")
    msgs = [{"ts": "1.0", "user": "UBOT", "bot_id": "BME", "text": "취합"},  # 실행기 자신
            {"ts": "2.0", "user": "U1", "text": "",  # 내려받기 실패: 오류 답글 후 다음 메시지 진행
             "files": [{"id": "F404", "name": "OLA_DCAMP9_S02_X_v0.1.0.pptx"}]},
            {"ts": "3.0", "user": "U1", "text": "",
             "files": [{"id": "F1", "name": "OLA_DCAMP9_S01_X_v0.1.0.pptx"}]},
            {"ts": "4.0", "bot_id": "BWF", "text": "승인 S01 v0.1.0"},  # 봇 승인은 받지 않음(4-Eyes)
            {"ts": "5.0", "bot_id": "BWF", "text": "취합 초안", "thread_ts": "3.0"}]  # 취합 버튼
    monkeypatch.setattr(run, "to_pdf", lambda p, d: p)
    monkeypatch.setattr(run, "PDF_READY", True)
    sl, reg = FakeSlack(msgs, {"F1": s01.read_bytes()}), Registry(cursor="0.5")
    assert run.process(sl, reg, frozenset({"UBOT", "BME"})) == (True, False)
    assert reg.cursor == "5.0" and any("오류" in p for p in sl.posts)
    assert reg.find("S01", "0.1.0").status == "초안" and reg.find("S02", "0.1.0") is None
    assert len(reg.builds) == 1 and reg.builds[0].result == "성공" and reg.builds[0].requested_by == "BWF"


def test_main_keeps_state_in_slack_and_resumes_after_pdf_install(tmp_path, monkeypatch):
    s01 = _slide(tmp_path / "S01.pptx", "S01")
    sl = FakeSlack([], {"F1": s01.read_bytes()})
    state, out = tmp_path / "work" / "state.json", tmp_path / "gh_output"
    monkeypatch.setattr(run, "Slack", lambda token: sl)
    monkeypatch.setattr(run, "to_pdf", lambda p, d: p)
    for k, v in {"SLACK_BOT_TOKEN": "x", "IR_STATE_PATH": str(state), "IR_START_TS": "0.5",
                 "GITHUB_OUTPUT": str(out), "IR_MAIN_CANVAS": "FMAIN"}.items():
        monkeypatch.setenv(k, v)
    monkeypatch.delenv("IR_BOOTSTRAP", raising=False)
    assert run.main() == 1 and not sl.private  # 상태 파일이 없으면 bootstrap 없이 시작하지 않는다

    monkeypatch.setenv("IR_BOOTSTRAP", "1")
    assert run.main() == 0
    assert list(sl.private) == ["FU1"] and "현황 캔버스" in sl.posts[0]

    # 다음 예약 실행: 새 러너라 작업 사본이 없고, Slack의 상태 파일에서 이어 받는다
    monkeypatch.delenv("IR_BOOTSTRAP")
    state.unlink()
    sl.messages += [{"ts": "1.0", "user": "U1", "text": "",
                     "files": [{"id": "F1", "name": "OLA_DCAMP9_S01_X_v0.1.0.pptx"}]},
                    {"ts": "2.0", "user": "U2", "text": "취합 초안"}]
    monkeypatch.setattr(run, "PDF_READY", False)
    assert run.main() == 0
    assert "needs_pdf=true" in out.read_text()  # 취합 명령 앞에서 멈추고 LibreOffice 설치를 요청
    reg = Registry.load(state)
    assert reg.cursor == "1.0" and len(reg.versions) == 1 and not reg.builds
    assert list(sl.private) == ["FU2"]  # 바뀐 상태를 새로 올리고 이전 상태 파일은 지운다

    # 같은 실행의 두 번째 단계: 작업 사본에서 이어서 취합한다
    monkeypatch.setattr(run, "PDF_READY", True)
    assert run.main() == 0
    reg = Registry.load(state)
    assert reg.cursor == "2.0" and reg.builds[-1].result == "성공"
    assert sl.canvases["FMAIN"].startswith("최신 취합본: **") and "최신 취합본: **" in sl.canvases["FC"]
    assert len(sl.private) == 1


def test_new_messages_picks_new_replies_in_older_threads_once():
    class S(Slack):
        def _messages(self, method, **p):
            if method == "conversations.history":
                return [{"ts": "200.0"},
                        {"ts": "150.0", "thread_ts": "50.0", "subtype": "thread_broadcast"},
                        {"ts": "50.0", "reply_count": 3, "latest_reply": "150.0"},
                        {"ts": "40.0", "reply_count": 1, "latest_reply": "45.0"}]
            assert method == "conversations.replies" and p["ts"] == "50.0"  # 새 답글이 있는 스레드만 연다
            return [{"ts": "50.0"}, {"ts": "60.0"}, {"ts": "120.0"}, {"ts": "150.0", "thread_ts": "50.0"}]

    assert [m["ts"] for m in S("x").new_messages("C", "100.0")] == ["120.0", "150.0", "200.0"]


def test_pdf_conversion_when_libreoffice_available(tmp_path):
    """CI(IR_REQUIRE_PDF=1)에서는 LibreOffice로 실제 PDF 장수까지 확인한다."""
    import os
    import subprocess

    from ir_deck.merge import to_pdf

    out = tmp_path / "out.pptx"
    merge([_slide(tmp_path / "a.pptx", "S01"), _slide(tmp_path / "b.pptx", "S02", chart=True),
           _slide(tmp_path / "c.pptx", "S03", theme_tag="T2")], out)
    try:
        pdf = to_pdf(out, tmp_path)
    except (MergeError, OSError, subprocess.SubprocessError):
        if os.environ.get("IR_REQUIRE_PDF"):
            raise
        pytest.skip("LibreOffice 변환 불가 환경")
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    pages = re.search(r"^Pages:\s+(\d+)", info, re.M)
    assert pages and int(pages.group(1)) == 3, info
